import openpyxl
import os
import json

files = [
    r'C:\Users\Delkira\Desktop\RELACION VACACIONES RRHH ACTUAL FEBRERO 2025 (Autoguardado).xlsx',
    r'C:\Users\Delkira\Desktop\RELACION VACACIONES RRHH ACTUAL FEBRERO 2025.xlsx',
    r'C:\Users\Delkira\Desktop\RELACION VACACIONES RRHH ACTUAL FEBRERO 20253.xlsx',
    r'C:\Users\Delkira\Desktop\VACACIONES RRHH\RELACION VACACIONES RRHH ACTUAL FEBRERO 2025 (Autoguardado).xlsx',
    r'C:\Users\Delkira\Desktop\VACACIONES RRHH\RELACION VACACIONES RRHH ACTUAL FEBRERO 2025.xlsx'
]

all_workers = {}

for f in files:
    if not os.path.exists(f):
        print(f"File not found: {f}")
        continue

    print(f"\nProcessing: {f}")
    wb = openpyxl.load_workbook(f, data_only=True)

    # Process 'vacaciones' sheet
    if 'vacaciones' in wb.sheetnames:
        ws = wb['vacaciones']
        for r in range(2, ws.max_row + 1):
            row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if not any(row): continue

            cargo = str(row[0] or '').strip()
            apellido = str(row[1] or '').strip()
            nombre = str(row[2] or '').strip()
            cedula_val = str(row[3] or '').strip().replace('.', '').replace('V-', '').replace('E-', '')
            if not cedula_val.isdigit(): continue
            cedula = f"V-{int(cedula_val):,}".replace(',', '.')
            full_name = f"{nombre} {apellido}".strip().title()

            # Read all period pairs starting from col 9 (index 8)
            periods = []
            c = 8
            while c < len(row) - 1:
                p_lbl = str(row[c] or '').strip()
                p_val = str(row[c+1] or '').strip()
                if p_lbl != 'None' and p_lbl != '' or p_val != 'None' and p_val != '':
                    periods.append({"period": p_lbl, "dates_or_status": p_val})
                c += 2

            if cedula not in all_workers:
                all_workers[cedula] = {
                    "cedula": cedula,
                    "name": full_name,
                    "cargo": cargo,
                    "periods": periods
                }
            else:
                # Merge unique periods
                existing_p = {p["period"]: p["dates_or_status"] for p in all_workers[cedula]["periods"]}
                for p in periods:
                    lbl = p["period"]
                    val = p["dates_or_status"]
                    if lbl not in existing_p or (existing_p[lbl] in ['None', '', 'None/None'] and val not in ['None', '']):
                        existing_p[lbl] = val
                all_workers[cedula]["periods"] = [{"period": k, "dates_or_status": v} for k, v in existing_p.items()]

print(f"\nTotal unique workers extracted: {len(all_workers)}")

# Filter workers with pending vacations
pending_workers = []
for ci, data in all_workers.items():
    # A period is pending if dates_or_status is empty, None, None/None or 'None'
    pending_periods = []
    for p in data["periods"]:
        st = p["dates_or_status"]
        if not st or st in ['None', 'None/None', ''] or 'None' in st:
            pending_periods.append(p["period"])

    if pending_periods or not data["periods"]:
        pending_workers.append({
            "cedula": data["cedula"],
            "name": data["name"],
            "cargo": data["cargo"],
            "all_periods": data["periods"],
            "pending_periods": pending_periods
        })

print(f"Workers with pending vacations: {len(pending_workers)}")

with open(r'C:\Users\Delkira\Downloads\IVSS\vacaciones_merged.json', 'w', encoding='utf-8') as jf:
    json.dump({"all": list(all_workers.values()), "pending": pending_workers}, jf, ensure_ascii=False, indent=2)

print("Saved merged data to vacaciones_merged.json")
